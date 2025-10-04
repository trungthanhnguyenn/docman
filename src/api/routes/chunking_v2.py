"""
Chunking V2 Routes - Advanced document chunking endpoints

Provides intelligent document chunking with multiple strategies:
- Content-based chunking (POST /v2/chunk/content)
- Document-based chunking (POST /v2/chunk/document/{document_id})

This is separate from the original chunks routes to maintain backward compatibility.
"""

import logging
from typing import Dict, Any

from fastapi import APIRouter, HTTPException, Depends
from datetime import datetime

from src.core.models import (
    ChunkingRequest,
    ChunkingV2Response,
    ChunkV2,
    DocumentChunkingRequest
)
from src.api.services.chunking_service import (
    get_chunking_service,
    ChunkingConfig
)
from src.api.services.database_manager import DatabaseManager
from src.api.dependencies import get_database_manager
from src.core.exceptions import DatabaseConnectionException

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/v2/chunk", tags=["chunking-v2"])


@router.post("/content", response_model=ChunkingV2Response)
async def chunk_content(request: ChunkingRequest) -> ChunkingV2Response:
    """
    Chunk document content with intelligent strategy selection.
    
    This endpoint accepts raw document content and returns intelligently chunked segments.
    
    **Chunking Modes:**
    - `auto`: Automatically detect best strategy based on content structure
    - `header_based`: Chunk by document headers (h1, h2, h3, etc.)
    - `qa_based`: Chunk by Q&A pairs (questions and answers)
    - `semantic`: Chunk by semantic boundaries (paragraphs)
    - `size_based`: Fixed-size chunks with overlap
    - `advanced_semantic`: Use ML models for semantic clustering (requires sentence-transformers)
    
    **Example Usage:**
    ```json
    {
        "content": "Your document text here...",
        "chunk_mode": "auto",
        "chunk_size": 1000,
        "chunk_overlap": 200
    }
    ```
    
    Args:
        request: ChunkingRequest with content and configuration
        
    Returns:
        ChunkingV2Response: Chunks with metadata
        
    Raises:
        HTTPException: If chunking fails
    """
    try:
        start_time = datetime.utcnow()
        
        # Create chunking config
        config = ChunkingConfig(
            chunk_size=request.chunk_size or 1000,
            chunk_overlap=request.chunk_overlap or 200,
            min_chunk_size=request.min_chunk_size or 50
        )
        
        # Get chunking service
        chunking_service = get_chunking_service(config)
        
        # Perform chunking
        logger.info(
            f"Chunking content (mode={request.chunk_mode}, "
            f"size={len(request.content)}, doc_id={request.document_id})"
        )
        
        result = chunking_service.chunk_content(
            content=request.content,
            chunk_mode=request.chunk_mode,
            filename=request.filename,
            document_id=request.document_id
        )
        
        # Convert to response model
        chunks_v2 = [
            ChunkV2(
                content=chunk.content,
                chunk_index=chunk.chunk_index,
                start_char=chunk.start_char,
                end_char=chunk.end_char,
                chunk_type=chunk.chunk_type,
                metadata=chunk.metadata,
                token_count=chunk.token_count,
                language=chunk.language
            )
            for chunk in result.chunks
        ]
        
        logger.info(
            f"Chunking completed: {result.total_chunks} chunks, "
            f"strategy={result.strategy_used}, time={result.processing_time_ms}ms"
        )
        
        return ChunkingV2Response(
            chunks=chunks_v2,
            total_chunks=result.total_chunks,
            strategy_used=result.strategy_used,
            processing_time_ms=result.processing_time_ms,
            document_metadata=result.document_metadata
        )
        
    except ValueError as e:
        logger.error(f"Validation error in chunking: {e}")
        raise HTTPException(
            status_code=400,
            detail=f"Chunking validation error: {str(e)}"
        )
    except Exception as e:
        logger.error(f"Chunking failed: {e}", exc_info=True)
        raise HTTPException(
            status_code=500,
            detail=f"Chunking failed: {str(e)}"
        )


@router.post("/document/{document_id}", response_model=ChunkingV2Response)
async def chunk_document_by_id(
    document_id: str,
    request: DocumentChunkingRequest,
    db_manager: DatabaseManager = Depends(get_database_manager)
) -> ChunkingV2Response:
    """
    Chunk a document by its ID.
    
    This endpoint:
    1. Downloads the document from MinIO using document_id
    2. Extracts text content from the document
    3. Chunks the content using the specified strategy
    
    **Workflow:**
    1. Upload document → `POST /api/v1/documents/session/{session_id}/upload`
    2. Get document_id from response
    3. Chunk document → `POST /api/v1/v2/chunk/document/{document_id}`
    4. Upload chunks with embeddings → `POST /api/v1/chunks/session/{session_id}/chunks`
    
    **Supported File Types:**
    - PDF (.pdf) - Extracts text using PyPDF2
    - Word (.docx) - Extracts text and tables using python-docx
    - Text (.txt, .md) - Direct text extraction
    
    Args:
        document_id: Document identifier
        request: DocumentChunkingRequest with chunking configuration
        db_manager: Database manager instance
        
    Returns:
        ChunkingV2Response: Chunks with metadata
        
    Raises:
        HTTPException: If document not found or chunking fails
    """
    try:
        # Validate document_id matches request
        if request.document_id != document_id:
            raise HTTPException(
                status_code=400,
                detail="document_id in path and request body must match"
            )
        
        logger.info(f"Chunking document {document_id} with mode={request.chunk_mode}")
        
        # Step 1: Download document from MinIO
        download_result = await db_manager.download_document(document_id=document_id)
        
        if download_result.get("error"):
            if "not found" in download_result["error"].lower():
                raise HTTPException(
                    status_code=404,
                    detail=f"Document not found: {document_id}"
                )
            else:
                raise HTTPException(
                    status_code=500,
                    detail=f"Failed to download document: {download_result['error']}"
                )
        
        file_content = download_result["file_content"]
        filename = download_result["filename"]
        content_type = download_result["content_type"]
        
        logger.info(
            f"Downloaded document: {filename} "
            f"(type={content_type}, size={len(file_content)} bytes)"
        )
        
        # Step 2: Extract text content based on file type
        try:
            text_content = _extract_text_content(
                file_content=file_content,
                filename=filename,
                content_type=content_type
            )
        except Exception as e:
            logger.error(f"Text extraction failed for {filename}: {e}")
            raise HTTPException(
                status_code=422,
                detail=f"Failed to extract text from document: {str(e)}"
            )
        
        if not text_content or not text_content.strip():
            raise HTTPException(
                status_code=422,
                detail=f"No text content could be extracted from document {filename}"
            )
        
        logger.info(f"Extracted text: {len(text_content)} characters")
        
        # Step 3: Chunk the content
        config = ChunkingConfig(
            chunk_size=request.chunk_size or 1000,
            chunk_overlap=request.chunk_overlap or 200
        )
        
        chunking_service = get_chunking_service(config)
        
        result = chunking_service.chunk_content(
            content=text_content,
            chunk_mode=request.chunk_mode,
            filename=filename,
            document_id=document_id
        )
        
        # Convert to response model
        chunks_v2 = [
            ChunkV2(
                content=chunk.content,
                chunk_index=chunk.chunk_index,
                start_char=chunk.start_char,
                end_char=chunk.end_char,
                chunk_type=chunk.chunk_type,
                metadata={
                    **chunk.metadata,
                    "document_id": document_id,
                    "filename": filename,
                    "content_type": content_type
                },
                token_count=chunk.token_count,
                language=chunk.language
            )
            for chunk in result.chunks
        ]
        
        logger.info(
            f"Document chunking completed: {result.total_chunks} chunks, "
            f"strategy={result.strategy_used}, time={result.processing_time_ms}ms"
        )
        
        return ChunkingV2Response(
            chunks=chunks_v2,
            total_chunks=result.total_chunks,
            strategy_used=result.strategy_used,
            processing_time_ms=result.processing_time_ms,
            document_metadata={
                **result.document_metadata,
                "document_id": document_id,
                "filename": filename,
                "content_type": content_type,
                "original_size_bytes": len(file_content),
                "extracted_text_length": len(text_content)
            }
        )
        
    except HTTPException:
        raise
    except DatabaseConnectionException as e:
        logger.error(f"Database connection error: {e}")
        raise HTTPException(
            status_code=503,
            detail=f"Database connection error: {e.message}"
        )
    except Exception as e:
        logger.error(f"Document chunking failed: {e}", exc_info=True)
        raise HTTPException(
            status_code=500,
            detail=f"Document chunking failed: {str(e)}"
        )


def _extract_text_content(
    file_content: bytes,
    filename: str,
    content_type: str
) -> str:
    """
    Extract text content from file based on type.
    
    Args:
        file_content: Raw file bytes
        filename: Original filename
        content_type: MIME type or file extension
        
    Returns:
        Extracted text content
        
    Raises:
        ValueError: If file type not supported or extraction fails
    """
    import io
    
    # Determine file type
    file_type = content_type.lower()
    if not file_type or file_type == "unknown":
        # Fallback to file extension
        file_type = filename.split('.')[-1].lower() if '.' in filename else 'txt'
    
    logger.info(f"Extracting text from file type: {file_type}")
    
    try:
        # PDF extraction
        if file_type in ['pdf', 'application/pdf']:
            try:
                import PyPDF2
            except ImportError:
                raise ValueError("PyPDF2 not installed. Cannot extract PDF content.")
            
            pdf_file = io.BytesIO(file_content)
            pdf_reader = PyPDF2.PdfReader(pdf_file)
            
            text_parts = []
            for page in pdf_reader.pages:
                text = page.extract_text()
                if text:
                    text_parts.append(text)
            
            return "\n".join(text_parts)
        
        # DOCX extraction
        elif file_type in ['docx', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document']:
            try:
                from docx import Document
            except ImportError:
                raise ValueError("python-docx not installed. Cannot extract DOCX content.")
            
            docx_file = io.BytesIO(file_content)
            doc = Document(docx_file)
            
            text_parts = []
            
            # Extract paragraphs
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text)
            
            # Extract tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        text_parts.append(row_text)
            
            return "\n".join(text_parts)
        
        # Text files
        elif file_type in ['txt', 'md', 'text/plain', 'text/markdown']:
            # Try UTF-8 first
            try:
                return file_content.decode('utf-8')
            except UnicodeDecodeError:
                # Fallback to latin-1
                return file_content.decode('latin-1', errors='ignore')
        
        else:
            raise ValueError(
                f"Unsupported file type: {file_type}. "
                f"Supported types: pdf, docx, txt, md"
            )
    
    except Exception as e:
        logger.error(f"Text extraction error: {e}")
        raise ValueError(f"Failed to extract text: {str(e)}")


@router.get("/strategies", response_model=Dict[str, Any])
async def get_chunking_strategies() -> Dict[str, Any]:
    """
    Get available chunking strategies and their descriptions.
    
    Returns information about all available chunking modes and when to use them.
    
    Returns:
        Dict with strategy information
    """
    return {
        "strategies": {
            "auto": {
                "name": "Automatic Strategy Selection",
                "description": "Automatically detects the best chunking strategy based on document structure",
                "use_when": "You're unsure which strategy to use or want optimal results",
                "available": True
            },
            "header_based": {
                "name": "Header-Based Chunking",
                "description": "Chunks document by headers (H1, H2, H3, numbered sections)",
                "use_when": "Document has clear header structure (e.g., manuals, reports)",
                "available": True
            },
            "qa_based": {
                "name": "Q&A-Based Chunking",
                "description": "Chunks document by question-answer pairs",
                "use_when": "Document is in Q&A format (e.g., FAQs, interviews)",
                "available": True
            },
            "semantic": {
                "name": "Semantic Boundary Chunking",
                "description": "Chunks by semantic boundaries (paragraphs and natural breaks)",
                "use_when": "Document has clear paragraph structure",
                "available": True
            },
            "size_based": {
                "name": "Fixed-Size Chunking",
                "description": "Creates fixed-size chunks with configurable overlap",
                "use_when": "Simple, uniform chunking needed or as fallback",
                "available": True
            },
            "advanced_semantic": {
                "name": "Advanced Semantic Clustering",
                "description": "Uses ML models (sentence-transformers) for semantic clustering",
                "use_when": "Need high-quality semantic chunks and have ML libraries installed",
                "available": SENTENCE_TRANSFORMERS_AVAILABLE
            }
        },
        "default_config": {
            "chunk_size": 1000,
            "chunk_overlap": 200,
            "min_chunk_size": 50,
            "max_chunk_size": 2000
        },
        "dependencies": {
            "sentence_transformers": SENTENCE_TRANSFORMERS_AVAILABLE,
            "langdetect": LANGDETECT_AVAILABLE
        }
    }


# Import availability flags
try:
    from sentence_transformers import SentenceTransformer
    SENTENCE_TRANSFORMERS_AVAILABLE = True
except ImportError:
    SENTENCE_TRANSFORMERS_AVAILABLE = False

try:
    from langdetect import detect
    LANGDETECT_AVAILABLE = True
except ImportError:
    LANGDETECT_AVAILABLE = False
