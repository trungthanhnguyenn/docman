# src/core/utils.py
import hashlib
import uuid
import mimetypes
from typing import Optional, List, Dict, Any, Union
from datetime import datetime
import logging
import requests
from io import BytesIO
from pathlib import Path
from docx import Document as DocxDocument
import os
from dotenv import load_dotenv

# Load environment variables
load_dotenv()
API_BASE_URL = os.getenv("API_BASE_URL", "http://localhost:8000")

logger = logging.getLogger(__name__)

def generate_document_id(filename: str, content: Optional[bytes] = None) -> str:
    """
    Generate a unique document ID based on filename and content.
    
    Args:
        filename: Original filename
        content: File content bytes (optional)
        
    Returns:
        str: Unique document identifier
    """
    if content:
        # Use content hash for better uniqueness
        content_hash = hashlib.sha256(content).hexdigest()
        return content_hash[:32]  # Use first 32 characters
    else:
        # Fallback to filename-based hash with timestamp
        timestamp = str(int(datetime.utcnow().timestamp() * 1000))
        combined = f"{filename}_{timestamp}"
        return hashlib.md5(combined.encode()).hexdigest()

def generate_chunk_id(document_id: str, page: int, chunk_index: int) -> str:
    """
    Generate a unique chunk ID.
    
    Args:
        document_id: Document identifier
        page: Page number
        chunk_index: Chunk index within the page
        
    Returns:
        str: Unique chunk identifier
    """
    combined = f"{document_id}_{page}_{chunk_index}"
    return hashlib.md5(combined.encode()).hexdigest()

def calculate_file_hash(content: bytes, algorithm: str = "sha256") -> str:
    """
    Calculate file hash.
    
    Args:
        content: File content bytes
        algorithm: Hash algorithm (sha256, md5, sha1)
        
    Returns:
        str: File hash
    """
    if algorithm == "sha256":
        return hashlib.sha256(content).hexdigest()
    elif algorithm == "md5":
        return hashlib.md5(content).hexdigest()
    elif algorithm == "sha1":
        return hashlib.sha1(content).hexdigest()
    else:
        raise ValueError(f"Unsupported hash algorithm: {algorithm}")

def detect_content_type(filename: str, content: Optional[bytes] = None) -> str:
    """
    Detect content type from filename and content.
    
    Args:
        filename: File name
        content: File content bytes (optional)
        
    Returns:
        str: MIME type
    """
    # Try to guess from filename
    content_type, _ = mimetypes.guess_type(filename)
    
    if content_type:
        return content_type
    
    # Fallback based on extension
    extension = filename.lower().split('.')[-1] if '.' in filename else ''
    
    extension_mapping = {
        'pdf': 'application/pdf',
        'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        'txt': 'text/plain',
        'md': 'text/markdown',
        'rtf': 'application/rtf'
    }
    
    return extension_mapping.get(extension, 'application/octet-stream')

def validate_file_type(filename: str, allowed_types: Optional[List[str]] = None) -> bool:
    """
    Validate if file type is allowed.
    
    Args:
        filename: File name
        allowed_types: List of allowed extensions
        
    Returns:
        bool: True if file type is allowed
    """
    if allowed_types is None:
        allowed_types = ['pdf', 'docx', 'txt', 'md', 'rtf']
    
    extension = filename.lower().split('.')[-1] if '.' in filename else ''
    return extension in allowed_types

def validate_file_size(file_size: int, max_size: Optional[int] = None) -> bool:
    """
    Validate if file size is within limits.
    
    Args:
        file_size: File size in bytes
        max_size: Maximum allowed size in bytes
        
    Returns:
        bool: True if file size is valid
    """
    if max_size is None:
        max_size = 50 * 1024 * 1024  # 50MB default
    
    return 0 < file_size <= max_size

def format_file_size(size_bytes: int) -> str:
    """
    Format file size in human-readable format.
    
    Args:
        size_bytes: Size in bytes
        
    Returns:
        str: Formatted size string
    """
    if size_bytes == 0:
        return "0 B"
    
    size_names = ["B", "KB", "MB", "GB", "TB"]
    size_float = float(size_bytes)
    i = 0
    while size_float >= 1024 and i < len(size_names) - 1:
        size_float /= 1024.0
        i += 1
    
    return f"{size_float:.1f} {size_names[i]}"

def sanitize_filename(filename: str) -> str:
    """
    Sanitize filename for safe storage.
    
    Args:
        filename: Original filename
        
    Returns:
        str: Sanitized filename
    """
    import re
    
    # Remove or replace dangerous characters
    filename = re.sub(r'[<>:"/\\|?*]', '_', filename)
    
    # Remove leading/trailing whitespace and dots
    filename = filename.strip(' .')
    
    # Ensure filename is not empty
    if not filename:
        filename = f"unnamed_{uuid.uuid4().hex[:8]}"
    
    return filename

def create_file_url(base_url: str, bucket_name: str, object_name: str) -> str:
    """
    Create a file URL for accessing stored documents.
    
    Args:
        base_url: Base URL of the storage service
        bucket_name: Storage bucket name
        object_name: Object name/key
        
    Returns:
        str: Complete file URL
    """
    base_url = base_url.rstrip('/')
    return f"{base_url}/{bucket_name}/{object_name}"

def parse_search_filters(filters: Dict[str, Any]) -> Dict[str, Any]:
    """
    Parse and validate search filters.
    
    Args:
        filters: Raw filter dictionary
        
    Returns:
        Dict[str, Any]: Parsed and validated filters
    """
    parsed_filters = {}
    
    # Handle document IDs
    if 'document_ids' in filters:
        doc_ids = filters['document_ids']
        if isinstance(doc_ids, str):
            parsed_filters['document_ids'] = [doc_ids]
        elif isinstance(doc_ids, list):
            parsed_filters['document_ids'] = doc_ids
    
    # Handle user IDs
    if 'user_ids' in filters:
        user_ids = filters['user_ids']
        if isinstance(user_ids, str):
            parsed_filters['user_ids'] = [user_ids]
        elif isinstance(user_ids, list):
            parsed_filters['user_ids'] = user_ids
    
    # Handle session IDs
    if 'session_ids' in filters:
        session_ids = filters['session_ids']
        if isinstance(session_ids, str):
            parsed_filters['session_ids'] = [session_ids]
        elif isinstance(session_ids, list):
            parsed_filters['session_ids'] = session_ids
    
    # Handle pages
    if 'pages' in filters:
        pages = filters['pages']
        if isinstance(pages, int):
            parsed_filters['pages'] = [pages]
        elif isinstance(pages, list):
            parsed_filters['pages'] = pages
    
    return parsed_filters

def chunk_text(text: str, chunk_size: int = 500, overlap: int = 50) -> List[str]:
    """
    Split text into overlapping chunks.
    
    Args:
        text: Text to chunk
        chunk_size: Maximum chunk size in characters
        overlap: Overlap between chunks
        
    Returns:
        List[str]: List of text chunks
    """
    if len(text) <= chunk_size:
        return [text]
    
    chunks = []
    start = 0
    
    while start < len(text):
        end = start + chunk_size
        
        # Try to break at word boundaries
        if end < len(text):
            # Find the last space before the end
            last_space = text.rfind(' ', start, end)
            if last_space > start:
                end = last_space
        
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        
        # Move start position with overlap
        start = end - overlap
        if start <= 0:
            start = end
    
    return chunks

def merge_metadata(existing: Dict[str, Any], new: Dict[str, Any]) -> Dict[str, Any]:
    """
    Merge metadata dictionaries with conflict resolution.
    
    Args:
        existing: Existing metadata
        new: New metadata to merge
        
    Returns:
        Dict[str, Any]: Merged metadata
    """
    merged = existing.copy()
    
    for key, value in new.items():
        if key not in merged:
            merged[key] = value
        elif isinstance(merged[key], dict) and isinstance(value, dict):
            merged[key] = merge_metadata(merged[key], value)
        else:
            # New value overwrites existing
            merged[key] = value
    
    return merged

def extract_text_from_docx(content_bytes: bytes) -> str:
    """Extract text content from .docx file bytes"""
    try:
        # Create a BytesIO object from bytes
        doc_stream = BytesIO(content_bytes)
        
        # Load the document
        doc = DocxDocument(doc_stream)
        
        # Extract text from all paragraphs
        text_content = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():  # Only add non-empty paragraphs
                text_content.append(paragraph.text.strip())
        
        # Join paragraphs with newlines
        full_text = '\n'.join(text_content)
        
        # If no text found, check tables
        if not full_text.strip():
            table_text = []
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_text.append(cell.text.strip())
                    if row_text:
                        table_text.append(' | '.join(row_text))
            
            if table_text:
                full_text = '\n'.join(table_text)
        
        return full_text if full_text.strip() else "No text content found in document"
        
    except Exception as e:
        return f"Error extracting text from docx: {str(e)}"

def get_documents_by_session(session_id: str, limit: int = 500, off_set: int = 0) -> List[Dict[str, Any]]:
    """Get documents for a specific session"""
    response = requests.get(f"{API_BASE_URL}/api/v1/sessions/{session_id}/documents", params={
        "limit": limit,
        "off_set": off_set
    })
    if response.status_code == 200:
        return response.json().get("documents", [])
    else:
        print(f"Failed to get documents: {response.text}")
        return []

def get_documents_content_by_indicators(session_id: str, indicators: List[Dict[str, str]]) -> List[Dict[str, str]]:
    """
    Get document contents by category and index indicators
    
    Args:
        session_id: Session identifier
        indicators: List of dicts with "category" and "index" keys
                   e.g. [{"category": "ami", "index": "1"}, {"category": "bci_1", "index": "1"}]
    
    Returns:
        List of dicts with "category" and "content" keys
        e.g. [{"category": "ami", "content": "Actual text content from file..."}, {"category": "bci_1", "content": "Actual text content..."}]
    """
    # Get all documents in the session
    documents = get_documents_by_session(session_id)
    
    if not documents:
        print(f"No documents found in session {session_id}")
        return []
    
    # Create lookup map from documents
    doc_lookup = {}
    for doc in documents:
        relative_path = doc.get('metadata', {}).get('relative_path', '')
        document_id = doc.get('document_id')
        filename = doc.get('filename', '')
        
        if relative_path and document_id:
            # Extract category and index from relative_path
            path_parts = relative_path.split('/')
            if len(path_parts) >= 2:
                folder_name = path_parts[-2]  # "ami", "bci_1", etc.
                filename_without_ext = path_parts[-1].split('.')[0]  # "ami_1", "bci_1", etc.
                
                if '_' in filename_without_ext:
                    file_prefix, file_index = filename_without_ext.split('_', 1)
                    
                    # Fixed logic for both patterns
                    # Pattern 1: ami/ami_1.docx -> key: ami_1 (category=ami, index=1) 
                    # Pattern 2: bci_3/bci_3.docx -> key: bci_3_3 (category=bci_3, index=3)
                    
                    if folder_name == file_prefix:
                        # Simple pattern: folder matches file prefix
                        key = f"{folder_name}_{file_index}"  # ami_1
                        category = folder_name  # ami
                        index = file_index      # 1
                    else:
                        # Numbered folder pattern: folder has number, file might have different number
                        key = f"{folder_name}_{file_index}"  # bci_3_3 
                        category = folder_name  # bci_3
                        index = file_index      # 3
                    
                    doc_lookup[key] = {
                        'document_id': document_id,
                        'category': category,
                        'index': index,
                        'filename': filename,
                        'relative_path': relative_path
                    }
    
    print(f"Created lookup map with {len(doc_lookup)} entries")
    
    # Find and download content for each indicator
    results = []
    
    for indicator in indicators:
        category = indicator.get('category')
        index = indicator.get('index')
        
        if not category or not index:
            results.append({
                'category': category or 'unknown',
                'content': f'Error: Invalid indicator {indicator}'
            })
            continue
        
        # Create lookup key
        lookup_key = f"{category}_{index}"
        
        if lookup_key in doc_lookup:
            doc_info = doc_lookup[lookup_key]
            document_id = doc_info['document_id']
            
            try:
                # Download document content
                print(f"Downloading {category}_{index} (ID: {document_id})...")
                response = requests.get(f"{API_BASE_URL}/api/v1/documents/{document_id}/download")
                
                if response.status_code == 200:
                    content_bytes = response.content
                    content_size = len(content_bytes)
                    
                    # Extract text based on file type
                    if doc_info['filename'].endswith('.txt'):
                        try:
                            content = content_bytes.decode('utf-8')
                        except UnicodeDecodeError:
                            try:
                                content = content_bytes.decode('latin-1')
                            except:
                                content = f"Error: Cannot decode text file ({content_size} bytes)"
                    
                    elif doc_info['filename'].endswith('.docx'):
                        # Extract text from docx
                        content = extract_text_from_docx(content_bytes)
                    
                    elif doc_info['filename'].endswith('.pdf'):
                        # For PDF files, you could use PyPDF2 or similar
                        content = f"PDF file detected ({content_size} bytes). Text extraction from PDF not implemented. Install PyPDF2 for PDF support."
                    
                    else:
                        # For other file types
                        content = f"Unsupported file type ({doc_info['filename']}, {content_size} bytes). Cannot extract text."
                    
                    results.append({
                        'category': category,
                        'content': content,
                        'document_id': document_id,
                        'filename': doc_info['filename'],
                        'relative_path': doc_info['relative_path'],
                        'file_size': content_size
                    })
                    
                    print(f"✓ Downloaded and extracted text from {category}_{index}: {len(content)} characters")
                    
                else:
                    results.append({
                        'category': category,
                        'content': f'Error downloading: HTTP {response.status_code}'
                    })
                    
            except Exception as e:
                results.append({
                    'category': category,
                    'content': f'Error downloading: {str(e)}'
                })
                
        else:
            # Document not found
            available_keys = list(doc_lookup.keys())[:10]  # Show first 10 available keys
            results.append({
                'category': category,
                'content': f'Document not found: {category}_{index}. Available keys: {available_keys}'
            })
            print(f"✗ Document {category}_{index} not found")
    
    return results

class Timer:
    """Simple timer utility for measuring execution time"""
    
    def __init__(self):
        self.start_time = None
        self.end_time = None
    
    def start(self):
        """Start the timer"""
        self.start_time = datetime.utcnow()
        return self
    
    def stop(self):
        """Stop the timer"""
        self.end_time = datetime.utcnow()
        return self
    
    def elapsed_ms(self) -> int:
        """Get elapsed time in milliseconds"""
        if self.start_time is None:
            return 0
        
        end = self.end_time or datetime.utcnow()
        delta = end - self.start_time
        return int(delta.total_seconds() * 1000)
    
    def __enter__(self):
        return self.start()
    
    def __exit__(self, exc_type, exc_val, exc_tb):
        self.stop()
